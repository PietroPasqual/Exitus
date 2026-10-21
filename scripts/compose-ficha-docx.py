"""Gera tests/fixtures/ficha-texto.docx a partir de tests/fixtures/ficha-texto.md.

Imita uma ficha digitada no Word: título no estilo "Title", marcadores pelos estilos de
lista ("List Bullet", "List Bullet 2"...), tabelas reais, tabulação entre recursos e
riscado de verdade. Um item recebe a numeração direto no parágrafo (<w:numPr>), como o
botão de marcadores do Word faz. Uso: python3 scripts/compose-ficha-docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'tests/fixtures/ficha-texto.md'
TARGET = ROOT / 'tests/fixtures/ficha-texto.docx'


def add_runs(paragraph, text):
    """Divide em runs nos trechos ~~riscados~~ e em pedaços menores, como o Word costuma salvar."""
    for part in re.split(r'(~~.+?~~)', text):
        if not part:
            continue
        struck = part.startswith('~~')
        part = part.strip('~')
        # O Word quebra o texto em vários runs (revisão, correção ortográfica...).
        for chunk in re.findall(r'.{1,7}', part, flags=re.S) if not struck else [part]:
            run = paragraph.add_run(chunk.replace('**', ''))
            run.font.strike = struck
            run.bold = '**' in text


def direct_bullet(paragraph, level):
    """Numeração no próprio parágrafo, sem estilo de lista."""
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), str(level))
    numid = OxmlElement('w:numId'); numid.set(qn('w:val'), '1')
    numpr.append(ilvl); numpr.append(numid); ppr.append(numpr)


def main():
    doc = Document()
    lines = SOURCE.read_text(encoding='utf8').split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-{3,}', c) or not c for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    add_runs(cell.paragraphs[0], value)
            continue
        level = (len(line) - len(line.lstrip())) // 4
        if stripped.startswith('# '):
            doc.add_paragraph(stripped[2:], style='Title')
        elif stripped.startswith('- '):
            text = stripped[2:]
            if text == 'Corda (2)':
                paragraph = doc.add_paragraph()
                direct_bullet(paragraph, level)
            else:
                paragraph = doc.add_paragraph(style='List Bullet' if level == 0 else f'List Bullet {min(level + 1, 3)}')
            add_runs(paragraph, text)
        else:
            paragraph = doc.add_paragraph()
            # Recursos lado a lado separados por tabulação, como no modelo.
            add_runs(paragraph, re.sub(r'\s{2,}(?=\S+: )', '\t', stripped))
        i += 1
    doc.save(TARGET)
    print(f'gerado {TARGET.relative_to(ROOT)}')


if __name__ == '__main__':
    main()
